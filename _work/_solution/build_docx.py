# -*- coding: utf-8 -*-
"""把 solution/产品解决方案说明书.md 渲染为带封面/目录/页眉页脚/截图的 DOCX。

用法（项目根目录）：
  <bundled-python> _work/_solution/build_docx.py solution/产品解决方案说明书.md solution/产品解决方案说明书.docx
"""
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))

# 复用项目既有的 Markdown → DOCX 转换器（本地工作区副本）
CONV_DIR = os.path.join(ROOT, '_work', '_archive-2026-10-25', 'scripts', 'misc')
sys.path.insert(0, CONV_DIR)
import md2docx  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402

NAVY = RGBColor(0x0A, 0x1B, 0x33)
NAVY2 = RGBColor(0x12, 0x30, 0x5A)
GOLD = RGBColor(0xC6, 0xA1, 0x5B)
GRAY = RGBColor(0x5C, 0x66, 0x75)
FONT_EA = '微软雅黑'
FONT_LATIN = 'Segoe UI'


def shade(paragraph, fill):
    pPr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def set_font(run, size=10.5, bold=False, color=None, ea=FONT_EA, latin=FONT_LATIN):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:eastAsia'), ea)


def new_para(doc, text='', size=10.5, bold=False, color=None, align=None,
             space_before=0, space_after=0, style=None, line_spacing=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line_spacing:
        pf.line_spacing = line_spacing
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def border(paragraph, color='C6A15B', size=8):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def field(paragraph, code, size=9, color=GRAY):
    run = paragraph.add_run()
    set_font(run, size=size, color=color)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = code
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def build_cover(doc):
    doc.add_paragraph()
    p = new_para(doc, 'CALR  ·  中国—东盟法律法规数据库', size=11, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=6, space_after=6)
    shade(p, '0A1B33')

    p = new_para(doc, '', space_before=40)
    p = new_para(doc, '产品解决方案说明书', size=30, bold=True, color=NAVY,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_before=40, space_after=10)
    p = new_para(doc, 'China–ASEAN Laws and Regulations Database', size=12,
                 color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p = new_para(doc, '可溯源 · 跨语言 · 可解释的东盟法律检索与 AI 问答平台', size=13,
                 color=NAVY2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    p = new_para(doc, '', space_after=8)
    border(p, 'C6A15B', 12)

    new_para(doc, '', space_before=30)
    new_para(doc, '2026 中国（广西）—东盟“AI+法律”应用创新大赛', size=12, bold=True,
             color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    new_para(doc, '高校学生赛道 · 面向东盟法律服务方向', size=11, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    new_para(doc, '上海对外经贸大学  ·  一粒橡子队', size=12, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    new_para(doc, '2026 年 10 月  ·  Version 1.0', size=10, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    new_para(doc, '', space_before=0)
    p = new_para(doc, '', space_before=20)
    border(p, '0A1B33', 4)
    page_break(doc)


def build_toc(doc, headings):
    p = new_para(doc, '目　录', size=20, bold=True, color=NAVY, space_after=8)
    border(p, 'C6A15B', 8)
    for text in headings:
        clean = re.sub(r'^\d+[.、]\s*', '', text).strip()
        new_para(doc, clean, size=11, color=NAVY2, space_before=1, space_after=3)
    new_para(doc, '（按章节顺序排列；页码见页脚）', size=9, color=GRAY, space_before=8)
    page_break(doc)


def add_header_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run('CALR 产品解决方案说明书')
    set_font(r, size=8.5, color=GRAY)

    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('第 ')
    set_font(r, size=9, color=GRAY)
    field(fp, 'PAGE')
    r = fp.add_run(' 页  /  共 ')
    set_font(r, size=9, color=GRAY)
    field(fp, 'NUMPAGES')
    r = fp.add_run(' 页')
    set_font(r, size=9, color=GRAY)


def replace_shots(doc):
    shots_dir = os.path.join(ROOT, 'solution', 'assets', 'shots')
    pattern = re.compile(r'^\[\[SHOT:([a-z0-9-]+)\|(.+)\]\]$')
    for p in list(doc.paragraphs):
        text = p.text.strip()
        m = pattern.match(text)
        if not m:
            continue
        name, caption = m.group(1), m.group(2).strip()
        img = os.path.join(shots_dir, name + '.png')
        if not os.path.exists(img):
            print('  [warn] missing shot:', img)
            p.text = caption
            continue
        p_el = p._p
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(6)
        img_p.paragraph_format.space_after = Pt(2)
        img_p.add_run().add_picture(img, width=Cm(15.6))
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_after = Pt(10)
        r = cap_p.add_run(caption)
        set_font(r, size=9, color=GRAY)
        p_el.addnext(img_p._p)
        img_p._p.addnext(cap_p._p)
        p_el.getparent().remove(p_el)


def main(md_path, docx_path):
    md_path = os.path.abspath(md_path)
    docx_path = os.path.abspath(docx_path)
    tmp = docx_path + '.base.docx'

    # 统一编码/换行：去 BOM，避免首个 "#" 被当成普通段落
    with open(md_path, 'r', encoding='utf-8-sig') as f:
        text = f.read().replace('\r\n', '\n')
    clean_md = docx_path + '.clean.md'
    with open(clean_md, 'w', encoding='utf-8', newline='\n') as f:
        f.write(text)
    md2docx.convert(clean_md, tmp)
    try:
        os.remove(clean_md)
    except OSError:
        pass

    doc = Document(tmp)
    body = doc.element.body
    title = '产品解决方案说明书'

    # 去掉 markdown 的第一个 H1（封面单独排版）
    for p in list(doc.paragraphs):
        if p.style is not None and p.style.name == 'Heading 1':
            if p.text.strip():
                title = p.text.strip()
            p._p.getparent().remove(p._p)
            break

    # 抽目录（H2）
    headings = []
    for p in doc.paragraphs:
        if p.style is not None and p.style.name in ('Heading 2', 'Heading 1'):
            t = p.text.strip()
            if t and t not in ('目录',):
                headings.append(t)

    # 先在文档末尾构建封面/目录，再整体移动到开头
    before = len(body)
    build_cover(doc)
    build_toc(doc, headings)
    after = len(body)
    new_nodes = list(body)[before:after]
    for node in reversed(new_nodes):
        body.remove(node)
        body.insert(0, node)

    add_header_footer(doc)
    replace_shots(doc)

    props = doc.core_properties
    props.title = title
    props.subject = '面向 12 法域的法律检索、跨语言比对与可溯源 AI 问答平台'
    props.author = '上海对外经贸大学 一粒橡子队'
    props.keywords = 'CALR; 中国—东盟; 法律法规数据库; 法律检索; 跨语言; AI 问答'
    props.comments = '产品解决方案说明书（比赛评审版）'

    doc.save(docx_path)
    try:
        os.remove(tmp)
    except OSError:
        pass
    print('OK', docx_path, os.path.getsize(docx_path), 'bytes')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
